import openpyxl
from openpyxl.styles import PatternFill
from pathlib import Path
from datetime import datetime
from collections import OrderedDict

# AOI-17→B(2), AOI-18→C(3), ..., AOI-25→J(10)
MACHINE_COL = {n: n - 17 + 2 for n in range(17, 26)}
ALL_DATA_COLS = list(range(2, 11))  # B~J

GRAY_FILL = PatternFill(start_color="FFD9D9D9", end_color="FFD9D9D9", fill_type="solid")
NO_FILL = PatternFill(fill_type=None)


def _parse_value(s):
    """'0.450000' → 0.45 / '2' → 2 / 'Surface' → 'Surface'"""
    try:
        return int(s)
    except (ValueError, TypeError):
        pass
    try:
        f = float(s)
        return int(f) if f == int(f) else f
    except (ValueError, TypeError):
        pass
    return s


def _values_equal(excel_val, ini_str):
    """엑셀 셀 값과 INI 문자열 값이 같은지 비교"""
    if excel_val is None:
        return False
    parsed = _parse_value(ini_str)
    try:
        return abs(float(excel_val) - float(parsed)) < 1e-9
    except (ValueError, TypeError):
        return str(excel_val).strip() == str(parsed).strip()


class ExcelHandler:
    def __init__(self, excel_path, machine_num):
        self.excel_path = Path(excel_path)
        self.machine_num = machine_num
        self.col = MACHINE_COL[machine_num]
        self.wb = openpyxl.load_workbook(excel_path)

    # ──────────────────────────────────────────────
    # 인덱스: A열 구조 파악
    # ──────────────────────────────────────────────
    def _build_index(self, ws):
        """
        Returns:
            section_rows: {section_name: row_number}
            param_rows:   {(section_name, param_name): row_number}
        """
        section_rows = {}
        param_rows = {}
        current_section = None

        for row in ws.iter_rows(min_row=2, max_col=1):
            cell = row[0]
            if cell.value is None:
                continue
            val = str(cell.value)
            if val.startswith('[') and val.endswith(']'):
                current_section = val
                section_rows[val] = cell.row
            elif current_section:
                param_rows[(current_section, val)] = cell.row

        return section_rows, param_rows

    def _find_insert_row(self, ws, section, section_rows):
        """섹션의 마지막 파라미터 다음 행 번호 반환"""
        if section not in section_rows:
            return ws.max_row + 1

        section_row = section_rows[section]
        last_row = section_row

        for row_num in range(section_row + 1, ws.max_row + 2):
            val = ws.cell(row=row_num, column=1).value
            if val is not None:
                if str(val).startswith('[') and str(val).endswith(']'):
                    break
                last_row = row_num

        return last_row + 1

    # ──────────────────────────────────────────────
    # 변경 사항 미리보기
    # ──────────────────────────────────────────────
    def preview_changes(self, parsed_files):
        """
        Returns list of change dicts:
          type: 'FILL'     — 빈칸에 값 채움
          type: 'UPDATE'   — 기존 값 변경
          type: 'NEW_ROW'  — 새 파라미터 행 추가
          type: 'NO_SHEET' — 시트 없음
        """
        changes = []

        for sheet_name, parsed in parsed_files.items():
            if sheet_name not in self.wb.sheetnames:
                changes.append({
                    'type': 'NO_SHEET', 'sheet': sheet_name,
                    'section': None, 'param': None,
                    'old_value': None, 'new_value': None, 'row': None,
                })
                continue

            ws = self.wb[sheet_name]
            section_rows, param_rows = self._build_index(ws)

            for section in parsed['sections_order']:
                for param, value in parsed['data'][section].items():
                    key = (section, param)
                    if key in param_rows:
                        row = param_rows[key]
                        old_val = ws.cell(row=row, column=self.col).value
                        if _values_equal(old_val, value):
                            continue  # 이미 동일
                        changes.append({
                            'type': 'FILL' if old_val is None else 'UPDATE',
                            'sheet': sheet_name, 'section': section,
                            'param': param,
                            'old_value': old_val,
                            'new_value': value,
                            'row': row,
                        })
                    else:
                        changes.append({
                            'type': 'NEW_ROW',
                            'sheet': sheet_name, 'section': section,
                            'param': param,
                            'old_value': None,
                            'new_value': value,
                            'row': None,
                        })

        return changes

    # ──────────────────────────────────────────────
    # 변경 사항 출력
    # ──────────────────────────────────────────────
    def print_changes(self, changes):
        from itertools import groupby

        by_sheet = OrderedDict()
        for c in changes:
            by_sheet.setdefault(c['sheet'], []).append(c)

        total_fill = sum(1 for c in changes if c['type'] == 'FILL')
        total_update = sum(1 for c in changes if c['type'] == 'UPDATE')
        total_new = sum(1 for c in changes if c['type'] == 'NEW_ROW')
        total_no_sheet = sum(1 for c in changes if c['type'] == 'NO_SHEET')

        for sheet_name, sheet_changes in by_sheet.items():
            fills   = [c for c in sheet_changes if c['type'] == 'FILL']
            updates = [c for c in sheet_changes if c['type'] == 'UPDATE']
            new_rows= [c for c in sheet_changes if c['type'] == 'NEW_ROW']
            no_sheet= [c for c in sheet_changes if c['type'] == 'NO_SHEET']

            print(f"\n  [{sheet_name}]")

            if no_sheet:
                print(f"    ⚠  시트 없음 — 건너뜀")
                continue

            if fills:
                print(f"    ▶ 채워넣기 ({len(fills)}개):")
                for c in fills:
                    print(f"       {c['param']:<45} → {c['new_value']}")

            if updates:
                print(f"    ▶ 값 변경 ({len(updates)}개):")
                for c in updates:
                    print(f"       {c['param']:<45} {c['old_value']}  →  {c['new_value']}")

            if new_rows:
                print(f"    ★ 새 파라미터 추가 ({len(new_rows)}개):")
                for c in new_rows:
                    print(f"       {c['section']} / {c['param']:<35} = {c['new_value']}")

        print()
        print(f"  요약: 채워넣기 {total_fill}개 | 값변경 {total_update}개 | 새행추가 {total_new}개", end='')
        if total_no_sheet:
            print(f" | 시트없음(건너뜀) {total_no_sheet}개", end='')
        print()

    # ──────────────────────────────────────────────
    # 변경 적용 및 저장
    # ──────────────────────────────────────────────
    def apply_changes(self, changes):
        # 1) FILL / UPDATE — 행 삽입 없으므로 먼저 처리
        for c in changes:
            if c['type'] in ('FILL', 'UPDATE'):
                ws = self.wb[c['sheet']]
                cell = ws.cell(row=c['row'], column=self.col)
                cell.value = _parse_value(c['new_value'])
                cell.fill = NO_FILL

        # 2) NEW_ROW — 시트별로 모아서 아래→위 순으로 삽입
        new_rows_by_sheet = OrderedDict()
        for c in changes:
            if c['type'] == 'NEW_ROW' and c['sheet'] in self.wb.sheetnames:
                new_rows_by_sheet.setdefault(c['sheet'], []).append(c)

        for sheet_name, new_rows in new_rows_by_sheet.items():
            ws = self.wb[sheet_name]
            section_rows, param_rows = self._build_index(ws)

            # 섹션별 묶기 (INI 순서 유지)
            section_params = OrderedDict()
            for c in new_rows:
                section_params.setdefault(c['section'], []).append(c)

            # 섹션별 삽입 위치 계산
            section_insert = {
                sec: self._find_insert_row(ws, sec, section_rows)
                for sec in section_params
            }

            # 섹션이 Excel에 없으면 끝에 헤더 추가
            for sec in section_params:
                if sec not in section_rows:
                    header_row = ws.max_row + 1
                    ws.cell(row=header_row, column=1).value = sec
                    section_rows[sec] = header_row
                    section_insert[sec] = header_row + 1

            # 아래→위 순으로 삽입 (행 번호 불변 보장)
            for sec in sorted(section_params, key=lambda s: section_insert[s], reverse=True):
                pos = section_insert[sec]
                for c in reversed(section_params[sec]):
                    ws.insert_rows(pos)
                    ws.cell(row=pos, column=1).value = c['param']
                    ws.cell(row=pos, column=self.col).value = _parse_value(c['new_value'])
                    for col in ALL_DATA_COLS:
                        if col != self.col:
                            ws.cell(row=pos, column=col).fill = GRAY_FILL

        # 저장 (타임스탬프 파일명)
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        out = self.excel_path.parent / f"{self.excel_path.stem}_AOI{self.machine_num}_{ts}{self.excel_path.suffix}"
        self.wb.save(out)
        return out
